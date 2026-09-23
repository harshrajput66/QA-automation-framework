from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pathlib import Path


GREEN_FILL = "C6EFCE"
RED_FILL = "FFC7CE"
HEADER_FILL = "D9EAF7"


def color_cell(cell, color):
    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")


def style_status_cell(cell, status):
    if status == "PASS":
        color_cell(cell, GREEN_FILL)
        cell.font = Font(bold=True, color="006100")
    else:
        color_cell(cell, RED_FILL)
        cell.font = Font(bold=True, color="9C0006")

    cell.alignment = Alignment(horizontal="center")


def generate_excel_report(sql_results, pandas_results, output_path="QA_Test_Execution_Report.xlsx"):
    print("Generating basic Excel QA report...")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "QA Report"

    headers = ["Test ID", "Expected Output", "Actual Output", "Status"]

    for column_number, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column_number)
        cell.value = header
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        color_cell(cell, HEADER_FILL)

    all_results = sql_results + pandas_results

    for row_number, result in enumerate(all_results, start=2):
        worksheet.cell(row=row_number, column=1).value = result["Test ID"]
        worksheet.cell(row=row_number, column=2).value = result.get("Expected", "")
        worksheet.cell(row=row_number, column=3).value = result.get("Result", "")
        worksheet.cell(row=row_number, column=4).value = result["Status"]

        for column_number in range(1, 4):
            worksheet.cell(row=row_number, column=column_number).alignment = Alignment(horizontal="center")

        style_status_cell(worksheet.cell(row=row_number, column=4), result["Status"])

    worksheet.column_dimensions["A"].width = 16
    worksheet.column_dimensions["B"].width = 22
    worksheet.column_dimensions["C"].width = 22
    worksheet.column_dimensions["D"].width = 14
    worksheet.freeze_panes = "A2"

    try:
        workbook.save(output_path)
        saved_path = output_path
    except PermissionError:
        path = Path(output_path)
        saved_path = str(path.with_name(path.stem + "_new" + path.suffix))
        workbook.save(saved_path)

    print(f"Basic Excel report saved: {saved_path}")
    return saved_path
