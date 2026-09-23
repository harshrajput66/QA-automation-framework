import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def add_title(document, text):
    heading = document.add_heading(text, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def add_heading(document, text):
    heading = document.add_heading(text, level=1)
    heading.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)


def add_bullet(document, bold_text, normal_text):
    paragraph = document.add_paragraph(style="List Bullet")
    bold_run = paragraph.add_run(bold_text + ": ")
    bold_run.bold = True
    paragraph.add_run(normal_text)


def set_default_font(document):
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def generate_doc(output_path="Project_Libraries_and_Terms.docx"):
    document = Document()
    set_default_font(document)

    add_title(document, "Data QA Project")
    add_title(document, "Interview Notes")

    intro = document.add_paragraph(
        "This document explains the project in simple language so it can be "
        "used for interview preparation."
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(document, "Project Summary")
    document.add_paragraph(
        "This project validates data after an ETL process. It reads sales data "
        "from a CSV file, creates source and target tables, runs SQL and Pandas "
        "checks, and creates Excel reports with PASS/FAIL results."
    )

    add_heading(document, "Pipeline")
    add_bullet(document, "Extract", "Read the Superstore CSV file using Pandas.")
    add_bullet(document, "Transform", "Create FactSales and dimension tables for a star schema.")
    add_bullet(document, "Load", "Save the source and target tables into SQLite.")
    add_bullet(document, "Validate", "Run SQL checks and Pandas/NumPy checks.")
    add_bullet(document, "Report", "Create Excel files for QA evidence.")

    add_heading(document, "Main Files")
    add_bullet(document, "main.py", "Runs the full pipeline from start to finish.")
    add_bullet(document, "kaggle_loader.py", "Loads data, builds tables, and adds demo defects.")
    add_bullet(document, "sql_validator.py", "Runs database checks using SQL.")
    add_bullet(document, "pandas_reconciler.py", "Compares source and target data with Pandas and NumPy.")
    add_bullet(document, "excel_reporter.py", "Creates the QA execution report.")
    add_bullet(document, "test_case_writer.py", "Creates the formal test case document.")

    add_heading(document, "Libraries")
    add_bullet(document, "pandas", "Used for CSV reading, data transformation, and DataFrame comparison.")
    add_bullet(document, "numpy", "Used for numeric tolerance checks with np.isclose().")
    add_bullet(document, "sqlite3", "Used to create and query a local SQLite database.")
    add_bullet(document, "openpyxl", "Used to create formatted Excel reports.")
    add_bullet(document, "python-docx", "Used to create this Word document.")

    add_heading(document, "Best Interview Explanation")
    document.add_paragraph(
        "I built a source-to-target reconciliation framework. The source is a "
        "Superstore CSV file. I transform it into a star schema and load it into "
        "SQLite. I then inject known target defects to prove the checks work. "
        "The framework validates nulls, duplicates, orphan records, row counts, "
        "aggregate totals, cell-level differences, and numeric tolerance. Finally, "
        "it generates Excel reports that can be shared with QA teams and business users."
    )

    document.save(output_path)
    print(f"Document saved: {output_path}")


if __name__ == "__main__":
    project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    output_file = os.path.join(project_root, "Project_Libraries_and_Terms.docx")
    generate_doc(output_path=output_file)
