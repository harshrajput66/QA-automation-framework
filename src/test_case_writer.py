from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from datetime import datetime
from pathlib import Path


def generate_test_cases_excel(sql_results, pandas_results, output_path="TestCases.xlsx"):
    print("Generating TestCases.xlsx...")

    test_case_info = {

        "TC-SQL-01": {
            "Objective":     "Check no NULL values exist in CustomerID, ProductID, LocationID in target table",
            "Precondition":  "ETL pipeline has loaded data into FactSales_Target in SQLite",
            "Input":         "FactSales_Target table - CustomerID, ProductID, LocationID columns"
        },
        "TC-SQL-02": {
            "Objective":     "Check no duplicate rows exist in target table based on OrderID + CustomerID + ProductID",
            "Precondition":  "FactSales_Target table is loaded by ETL pipeline",
            "Input":         "FactSales_Target - all rows grouped by OrderID, CustomerID, ProductID"
        },
        "TC-SQL-03": {
            "Objective":     "Check every CustomerID in FactSales exists in DimCustomer (referential integrity)",
            "Precondition":  "DimCustomer and FactSales_Target are both populated",
            "Input":         "LEFT JOIN FactSales_Target with DimCustomer on CustomerID"
        },
        "TC-SQL-04": {
            "Objective":     "Check every ProductID in FactSales exists in DimProduct (referential integrity)",
            "Precondition":  "DimProduct and FactSales_Target are both populated",
            "Input":         "LEFT JOIN FactSales_Target with DimProduct on ProductID"
        },
        "TC-SQL-05": {
            "Objective":     "Check no Sales values are negative (business rule: Sales must be >= 0)",
            "Precondition":  "FactSales_Target is loaded. Business rule: Sales >= 0",
            "Input":         "FactSales_Target - Sales column, all rows"
        },
        "TC-SQL-06": {
            "Objective":     "Check total row count in target matches source - no rows lost or added during ETL",
            "Precondition":  "Both FactSales_Source and FactSales_Target are loaded in SQLite",
            "Input":         "COUNT(*) from FactSales_Source vs COUNT(*) from FactSales_Target"
        },
        "TC-SQL-07": {
            "Objective":     "Check total Sales sum in target matches source - no data corruption during ETL",
            "Precondition":  "Both FactSales_Source and FactSales_Target are loaded in SQLite",
            "Input":         "SUM(Sales) from FactSales_Source vs SUM(Sales) from FactSales_Target"
        },
        "TC-PD-01": {
            "Objective":     "Check that column names of source and target DataFrames match",
            "Precondition":  "Source and target DataFrames are loaded in Python",
            "Input":         "source.columns vs target.columns using Pandas"
        },
        "TC-PD-02": {
            "Objective":     "Check that total rows in target DataFrame equals rows in source DataFrame",
            "Precondition":  "Source and target DataFrames are loaded in Python",
            "Input":         "len(source) vs len(target) using Pandas"
        },
        "TC-PD-03": {
            "Objective":     "Check no NULL values exist in OrderID, CustomerID, ProductID, LocationID, Sales",
            "Precondition":  "Target DataFrame is loaded in Python",
            "Input":         "target[critical_cols].isna().sum() using Pandas"
        },
        "TC-PD-04": {
            "Objective":     "Check no duplicate rows exist in target based on OrderID, CustomerID, ProductID",
            "Precondition":  "Target DataFrame is loaded in Python",
            "Input":         "target.duplicated(subset=[OrderID, CustomerID, ProductID]) using Pandas"
        },
        "TC-PD-05": {
            "Objective":     "Compare Sales, Quantity, Discount, Profit values cell by cell between source and target",
            "Precondition":  "Source and target DataFrames are aligned and loaded in Python",
            "Input":         "source[measure_cols].compare(target[measure_cols]) using Pandas df.compare()"
        },
        "TC-PD-06": {
            "Objective":     "Check Profit values are within 0.1% tolerance between source and target using NumPy",
            "Precondition":  "Source and target DataFrames are loaded. Tolerance: rtol=0.001, atol=0.01",
            "Input":         "numpy.isclose(source_profit, target_profit, rtol=0.001, atol=0.01)"
        },
    }

    all_results = sql_results + pandas_results
    test_cases  = []

    for result in all_results:
        test_id = result["Test ID"]
        info    = test_case_info.get(test_id, {})

        test_cases.append({
            "Test Case ID":    test_id,
            "Objective":       info.get("Objective",    ""),
            "Precondition":    info.get("Precondition", ""),
            "Input":           info.get("Input",        ""),
            "Expected Result": "Count of issues = 0 (no defects expected)",
            "Actual Result":   result.get("Details", ""),
            "Status":          result.get("Status", "")
        })

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    ws["A1"] = "Data QA - Test Cases Document"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A1:G1")
    ws.row_dimensions[1].height = 30

    passed = sum(1 for tc in test_cases if tc["Status"] == "PASS")
    failed = sum(1 for tc in test_cases if tc["Status"] == "FAIL")
    ws["A2"] = "Generated: " + datetime.now().strftime("%d-%b-%Y %H:%M") + \
               "   |   Total: " + str(len(test_cases)) + \
               "   |   PASS: " + str(passed) + \
               "   |   FAIL: " + str(failed)
    ws["A2"].font = Font(italic=True, size=10)
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 18

    ws.row_dimensions[3].height = 6

    columns = [
        "Test Case ID",
        "Objective",
        "Precondition",
        "Input",
        "Expected Result",
        "Actual Result",
        "Status"
    ]

    for col_num, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=4, column=col_num)
        cell.value     = col_name
        cell.font      = Font(bold=True, size=11, color="FFFFFF")
        cell.fill      = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[4].height = 25

    for row_num, tc in enumerate(test_cases, start=5):

        ws.cell(row=row_num, column=1).value = tc["Test Case ID"]
        ws.cell(row=row_num, column=2).value = tc["Objective"]
        ws.cell(row=row_num, column=3).value = tc["Precondition"]
        ws.cell(row=row_num, column=4).value = tc["Input"]
        ws.cell(row=row_num, column=5).value = tc["Expected Result"]
        ws.cell(row=row_num, column=6).value = tc["Actual Result"]
        ws.cell(row=row_num, column=7).value = tc["Status"]

        ws.row_dimensions[row_num].height = 55

        ws.cell(row=row_num, column=1).fill = PatternFill(
            start_color="DEEAF1", end_color="DEEAF1", fill_type="solid"
        )
        ws.cell(row=row_num, column=1).font = Font(bold=True, color="1F4E79")
        ws.cell(row=row_num, column=1).alignment = Alignment(horizontal="center", vertical="center")

        for col_num in range(2, 7):
            ws.cell(row=row_num, column=col_num).alignment = Alignment(
                horizontal="left", vertical="top", wrap_text=True
            )

        status_cell = ws.cell(row=row_num, column=7)
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        if tc["Status"] == "PASS":
            status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            status_cell.font = Font(bold=True, color="375623")
        else:
            status_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            status_cell.font = Font(bold=True, color="9C0006")

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 45
    ws.column_dimensions["E"].width = 35
    ws.column_dimensions["F"].width = 50
    ws.column_dimensions["G"].width = 10

    ws.freeze_panes = "A5"

    ws.auto_filter.ref = "A4:G4"

    try:
        wb.save(output_path)
        saved_path = output_path
    except PermissionError:
        path = Path(output_path)
        saved_path = str(path.with_name(path.stem + "_new" + path.suffix))
        wb.save(saved_path)

    print("TestCases.xlsx saved:", saved_path)
    return saved_path
